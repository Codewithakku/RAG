import os
import re
import csv

from typing import List

from openpyxl import load_workbook

from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    CSVLoader,
    UnstructuredMarkdownLoader,
)

from langchain_text_splitters import (
    RecursiveCharacterTextSplitter,
    NLTKTextSplitter,
)

from langchain_core.documents import Document

from app.config import CHUNK_SIZE, CHUNK_OVERLAP


class DocumentProcessor:
    """
    Handles loading and chunking of different document types.

    Supported formats:
        PDF  -> PyPDFLoader
        TXT  -> TextLoader
        MD   -> UnstructuredMarkdownLoader
        CSV  -> Row-Based Chunking
        XLSX -> Sheet + Row-Based Chunking

    Chunking strategies for text documents:
        recursive -> RecursiveCharacterTextSplitter
        sentence  -> NLTKTextSplitter
        paragraph -> Paragraph-based splitting

    CSV / XLSX:
        Row-based chunking is used automatically.
    """

    def __init__(
        self,
        chunk_size: int = CHUNK_SIZE,
        chunk_overlap: int = CHUNK_OVERLAP,
        chunking_strategy: str = "recursive"
    ):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.chunking_strategy = chunking_strategy.lower()

        if self.chunking_strategy == "recursive":

            self.text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=self.chunk_size,
                chunk_overlap=self.chunk_overlap,
                length_function=len,
                is_separator_regex=False,
            )

        elif self.chunking_strategy == "sentence":

            self.text_splitter = NLTKTextSplitter(
                chunk_size=self.chunk_size,
                chunk_overlap=self.chunk_overlap,
            )

        elif self.chunking_strategy == "paragraph":

            self.text_splitter = None

        else:

            raise ValueError(
                f"Unsupported chunking strategy: "
                f"{chunking_strategy}. "
                f"Supported strategies are: "
                f"recursive, sentence, paragraph"
            )

    # ============================================================
    # PDF Loader
    # ============================================================

    def load_pdf(
        self,
        file_path: str
    ) -> List[Document]:

        if not os.path.exists(file_path):

            raise FileNotFoundError(
                f"File not found at: {file_path}"
            )

        try:

            loader = PyPDFLoader(file_path)

            documents = loader.load()

            if documents and any(
                doc.page_content.strip()
                for doc in documents
            ):
                return documents

        except Exception as e:

            print(
                f"PyPDFLoader warning: {str(e)}. "
                "Trying direct pypdf fallback."
            )

        # --------------------------------------------------------
        # pypdf fallback
        # --------------------------------------------------------

        try:

            from pypdf import PdfReader

            reader = PdfReader(file_path)

            documents = []

            for page_num, page in enumerate(reader.pages):

                text = page.extract_text() or ""

                if text.strip():

                    documents.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": file_path,
                                "page": page_num
                            }
                        )
                    )

            return documents

        except Exception as fallback_err:

            raise ValueError(
                f"Failed to parse PDF document: "
                f"{str(fallback_err)}"
            )

    # ============================================================
    # TXT Loader
    # ============================================================

    def load_txt(
        self,
        file_path: str
    ) -> List[Document]:

        loader = TextLoader(
            file_path,
            encoding="utf-8"
        )

        return loader.load()

    # ============================================================
    # Markdown Loader
    # ============================================================

    def load_md(
        self,
        file_path: str
    ) -> List[Document]:

        loader = UnstructuredMarkdownLoader(
            file_path
        )

        return loader.load()

    # ============================================================
    # CSV Loader
    # ============================================================

    def load_csv(
        self,
        file_path: str
    ) -> List[Document]:

        loader = CSVLoader(
            file_path=file_path,
            encoding="utf-8"
        )

        return loader.load()

    # ============================================================
    # Excel Loader
    # ============================================================

    def load_excel(
        self,
        file_path: str
    ) -> List[Document]:

        # This loader is kept for compatibility.
        # Actual XLSX processing uses row-based processing.

        try:

            workbook = load_workbook(
                filename=file_path,
                read_only=True,
                data_only=True
            )

            workbook.close()

            return []

        except Exception as e:

            raise ValueError(
                f"Failed to read Excel file: {str(e)}"
            )

    # ============================================================
    # Generic Document Loader
    # ============================================================

    def load_document(
        self,
        file_path: str
    ) -> List[Document]:

        if not os.path.exists(file_path):

            raise FileNotFoundError(
                f"Document not found at: {file_path}"
            )

        extension = os.path.splitext(
            file_path
        )[1].lower()

        if extension == ".pdf":

            return self.load_pdf(file_path)

        elif extension == ".txt":

            return self.load_txt(file_path)

        elif extension == ".md":

            return self.load_md(file_path)

        elif extension == ".csv":

            return self.load_csv(file_path)

        elif extension == ".xlsx":

            return self.load_excel(file_path)

        else:

            raise ValueError(
                f"Unsupported file type: {extension}. "
                "Supported formats are: "
                ".pdf, .txt, .md, .csv, .xlsx"
            )

    # ============================================================
    # Paragraph-Based Chunking
    # ============================================================

    def paragraph_split_documents(
        self,
        documents: List[Document]
    ) -> List[Document]:

        chunks = []

        for doc in documents:

            text = doc.page_content

            text = text.replace("\r\n", "\n")
            text = text.replace("\r", "\n")

            paragraphs = re.split(
                r"\n\s*\n+",
                text
            )

            for paragraph in paragraphs:

                paragraph = paragraph.strip()

                if not paragraph:
                    continue

                paragraph = re.sub(
                    r"\s+",
                    " ",
                    paragraph
                )

                chunks.append(
                    Document(
                        page_content=paragraph,
                        metadata=doc.metadata.copy()
                    )
                )

        return chunks

    # ============================================================
    # CSV Row-Based Chunking
    # ============================================================

    def process_csv_rows(
        self,
        file_path: str
    ) -> List[Document]:

        chunks = []

        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            reader = csv.DictReader(file)

            for row_index, row in enumerate(
                reader,
                start=1
            ):

                row_values = []

                for key, value in row.items():

                    if value is None:
                        continue

                    value = str(value).strip()

                    if not value:
                        continue

                    row_values.append(
                        f"{key}: {value}"
                    )

                row_text = "\n".join(
                    row_values
                )

                if not row_text.strip():
                    continue

                chunks.append(
                    Document(
                        page_content=row_text,
                        metadata={
                            "source": file_path,
                            "row_index": row_index,
                            "chunk_type": "row"
                        }
                    )
                )

        return chunks

    # ============================================================
    # Excel Row-Based Chunking
    # ============================================================

    def process_excel_rows(
        self,
        file_path: str
    ) -> List[Document]:

        chunks = []

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True
        )

        for sheet in workbook.worksheets:

            rows = sheet.iter_rows(
                values_only=True
            )

            try:

                headers = next(rows)

            except StopIteration:

                continue

            headers = [
                str(header).strip()
                if header is not None
                else f"column_{i}"
                for i, header in enumerate(headers)
            ]

            for row_index, row in enumerate(
                rows,
                start=2
            ):

                row_values = []

                for i, value in enumerate(row):

                    if i >= len(headers):
                        continue

                    if value is None:
                        continue

                    value = str(value).strip()

                    if not value:
                        continue

                    row_values.append(
                        f"{headers[i]}: {value}"
                    )

                row_text = "\n".join(
                    row_values
                )

                if not row_text.strip():
                    continue

                chunks.append(
                    Document(
                        page_content=row_text,
                        metadata={
                            "source": file_path,
                            "sheet": sheet.title,
                            "row_index": row_index,
                            "chunk_type": "row"
                        }
                    )
                )

        workbook.close()

        return chunks

    # ============================================================
    # Process Document
    # ============================================================

    def process_document(
        self,
        file_path: str,
        doc_id: str
    ) -> List[Document]:

        if not os.path.exists(file_path):

            raise FileNotFoundError(
                f"Document not found at: {file_path}"
            )

        extension = os.path.splitext(
            file_path
        )[1].lower()

        file_name = os.path.basename(
            file_path
        )

        # ========================================================
        # CSV
        # ========================================================

        if extension == ".csv":

            chunks = self.process_csv_rows(
                file_path
            )

            chunking_strategy = "row"

        # ========================================================
        # XLSX
        # ========================================================

        elif extension == ".xlsx":

            chunks = self.process_excel_rows(
                file_path
            )

            chunking_strategy = "row"

        # ========================================================
        # PDF / TXT / MD
        # ========================================================

        else:

            raw_docs = self.load_document(
                file_path
            )

            if not raw_docs:

                return []

            # ----------------------------------------------------
            # Paragraph
            # ----------------------------------------------------

            if self.chunking_strategy == "paragraph":

                chunks = self.paragraph_split_documents(
                    raw_docs
                )

            # ----------------------------------------------------
            # Recursive / Sentence
            # ----------------------------------------------------

            else:

                chunks = self.text_splitter.split_documents(
                    raw_docs
                )

            chunking_strategy = (
                self.chunking_strategy
            )

        # ========================================================
        # Common Metadata
        # ========================================================

        for idx, chunk in enumerate(chunks):

            chunk.metadata["doc_id"] = doc_id

            chunk.metadata["source_file"] = file_name

            chunk.metadata["chunk_index"] = idx

            chunk.metadata["file_type"] = extension

            chunk.metadata["chunking_strategy"] = (
                chunking_strategy
            )

        return chunks